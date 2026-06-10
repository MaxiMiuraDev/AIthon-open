import openpyxl
import sys

f = sys.argv[1]
maxrows = int(sys.argv[2]) if len(sys.argv) > 2 else 80

wb = openpyxl.load_workbook(f, data_only=True)
print("Sheets:", wb.sheetnames)
for sheetname in wb.sheetnames:
    ws = wb[sheetname]
    print("\n=== Sheet:", sheetname, "dims:", ws.dimensions, "===")
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, maxrows), values_only=True):
        if any(c is not None for c in row):
            print(row)
