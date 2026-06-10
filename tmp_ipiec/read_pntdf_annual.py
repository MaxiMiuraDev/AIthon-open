import openpyxl

f = "16_1_04_PNTDF.xlsx"
wb = openpyxl.load_workbook(f, data_only=True)
ws = wb['16_1_04']

current_year = None
yearly_totals = {}
yearly_residentes = {}
yearly_no_residentes = {}
yearly_months_count = {}

for row in ws.iter_rows(min_row=4, max_row=ws.max_row, values_only=True):
    year_cell, month, total, _, residentes, no_residentes, _ = row[:7]
    if year_cell is not None:
        try:
            current_year = int(year_cell)
        except (ValueError, TypeError):
            continue
    if current_year is None:
        continue
    if month is None:
        continue

    def to_num(v):
        if v in ('-', '///', None):
            return 0
        try:
            return float(v)
        except (ValueError, TypeError):
            return 0

    yearly_totals[current_year] = yearly_totals.get(current_year, 0) + to_num(total)
    yearly_residentes[current_year] = yearly_residentes.get(current_year, 0) + to_num(residentes)
    yearly_no_residentes[current_year] = yearly_no_residentes.get(current_year, 0) + to_num(no_residentes)
    yearly_months_count[current_year] = yearly_months_count.get(current_year, 0) + 1

for year in sorted(yearly_totals.keys()):
    print(year, "months:", yearly_months_count[year], "total:", int(yearly_totals[year]),
          "residentes:", int(yearly_residentes[year]), "no_residentes:", int(yearly_no_residentes[year]))
