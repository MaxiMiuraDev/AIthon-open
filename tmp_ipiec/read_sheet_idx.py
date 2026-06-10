import openpyxl
import sys

f = sys.argv[1]
idx = int(sys.argv[2])
maxrows = int(sys.argv[3]) if len(sys.argv) > 3 else 100

wb = openpyxl.load_workbook(f, data_only=True)
sheetname = wb.sheetnames[idx]
ws = wb[sheetname]
print("=== Sheet index", idx, "name(repr):", repr(sheetname), "dims:", ws.dimensions, "===")
for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, maxrows), values_only=True):
    if any(c is not None for c in row):
        row2 = list(row)
        while row2 and row2[-1] is None:
            row2.pop()
        print(row2)
